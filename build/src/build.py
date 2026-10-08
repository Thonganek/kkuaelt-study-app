# -*- coding: utf-8 -*-
"""Assemble the standalone HTML study app from plan + content JSON + the learner's DOCX files.
Usage: python build.py [--out PATH] [--allow-missing]
"""
import json, re, sys, os, glob, datetime
from pathlib import Path
import docx

sys.stdout.reconfigure(encoding='utf-8')
HERE = Path(__file__).resolve().parent
SCR = HERE.parent
CONTENT = SCR / 'content'
SRC_DIR = Path(r'C:/Users/ACE/Desktop/kkuaelt')
sys.path.insert(0, str(HERE))
import plan_days as P
import plan42 as P42
import lessons as LS

args = sys.argv[1:]
OUT = Path(args[args.index('--out') + 1]) if '--out' in args else SCR / 'out' / 'KKU-AELT_42_Day_Study_Plan.html'
ALLOW_MISSING = '--allow-missing' in args
problems = []


def load(name, default=None):
    p = CONTENT / name
    if not p.exists():
        if ALLOW_MISSING:
            problems.append(f'missing {name}')
            return default
        raise SystemExit(f'missing content file {name}')
    return json.loads(p.read_text(encoding='utf-8'))


# ---------------- vocabulary ----------------
plan = json.loads((SCR / 'spec' / 'vocab_plan.json').read_text(encoding='utf-8'))
vocab = {}
for part in ('vocab_part1.json', 'vocab_part2.json'):
    v = load(part, {})
    for d, lst in (v or {}).items():
        vocab[d] = lst
for d, lst in plan.items():
    if d not in vocab:
        problems.append(f'vocab day {d} missing')
        vocab[d] = [{'w': x['w'], 'pos': '', 'th': '(ยังไม่มีความหมาย)', 'fam': '', 'col': '', 'ex': ''} for x in lst]
    assert [x['w'] for x in vocab[d]] == [x['w'] for x in lst], f'vocab order mismatch day {d}'
    for e, p in zip(vocab[d], lst):
        e['src'] = p['src']
vocab = {str(k): vocab[str(k)] for k in sorted(map(int, vocab))}
vocab = P42.remap_vocab(vocab)  # old 45-day numbering -> 42-day plan
assert sum(len(v) for v in vocab.values()) == 740

# ---------------- passages & tests ----------------
passages, tests = {}, {}
for name in ('weekday_A.json', 'weekday_B.json', 'weekday_C.json'):
    w = load(name, {'passages': []})
    for p in (w or {}).get('passages', []):
        passages[p['id']] = p
for tid in ('T1', 'T2', 'T3', 'T4', 'T5', 'T6', 'M1', 'M2', 'M3'):
    t = load(f'test_{tid}.json', None)
    if not t:
        continue
    ids = []
    for p in t['passages']:
        passages[p['id']] = p
        ids.append(p['id'])
    tests[tid] = {'id': tid, 'day': t['day'], 'title': t['title'], 'passages': ids}

# ---------------- grammar & essays ----------------
grammar = {g['id']: g for g in (load('grammar.json', {'sets': []}) or {}).get('sets', [])}
essays = {e['id']: e for e in (load('essays.json', {'essays': []}) or {}).get('essays', [])}

# ---------------- DOCX models ----------------
def paras(path):
    d = docx.Document(str(path))
    return [(p.style.name, p.text.strip()) for p in d.paragraphs if p.text.strip()], d.tables

models = {}
gp, gt = paras(SRC_DIR / 'Opinion_Essay_Guide.docx')
cur = None; n = 0
for style, text in gp:
    if style.startswith('Heading 2') and text.startswith('Essay'):
        n += 1; cur = {'title': text, 'source': 'Opinion_Essay_Guide.docx', 'topic': '', 'paras': [], 'notes': []}; models[f'GE{n}'] = cur; continue
    if style.startswith('Heading'):
        cur = None if not style.startswith('Heading 2') or not text.startswith('Essay') else cur; continue
    if cur is None:
        continue
    if text.startswith('Topic:'):
        cur['topic'] = text
    elif text.startswith('(Word count'):
        cur['notes'].append(text + ' — นับจริงด้วยโปรแกรมอาจไม่ตรงกับตัวเลขที่ระบุในเอกสาร')
    else:
        cur['paras'].append(text)
guide_words = [[c.text.strip() for c in r.cells] for r in gt[0].rows] if gt else []

ip, it = paras(SRC_DIR / 'Improved_Opinion_Essays.docx')
sections = []
for style, text in ip:
    if style.startswith('Heading 1'):
        sections.append({'h': text, 'lines': []})
    elif sections:
        sections[-1]['lines'].append((style, text))
k = 0
principles = None
ANALYSIS_START = ('เรื่องนี้มีปัญหา', 'วิเคราะห์ปัญหาของต้นฉบับ')
for sec in sections:
    if sec['h'].startswith('สรุปหลักการ'):
        principles = {'title': 'สรุปหลักการปรับปรุงสำคัญ (เพื่อให้ได้ 80+)', 'source': 'Improved_Opinion_Essays.docx', 'topic': '', 'paras': [t for s, t in sec['lines']], 'notes': []}
        continue
    if not sec['h'].startswith('Essay'):
        continue
    k += 1
    m = {'title': '', 'source': 'Improved_Opinion_Essays.docx', 'topic': '', 'stance': '', 'original': [], 'analysis': [], 'paras': [], 'notes': []}
    mode = None
    for style, text in sec['lines']:
        if text.startswith('Topic:'):
            m['topic'] = text; continue
        if text.startswith('Stance:'):
            m['stance'] = text; continue
        if text.startswith('(Word count'):
            continue
        if text == 'พี่เขียนเอง':
            mode = 'original'; continue
        if text.startswith(ANALYSIS_START):
            mode = 'analysis'
            if text.startswith('วิเคราะห์'):
                continue
        if text.startswith('Essay ฉบับปรับปรุง'):
            mode = 'pre'; continue
        if text.startswith('[') and text.endswith(']'):
            mode = 'improved'
        if text.startswith(('จุดที่ปรับปรุง', 'สิ่งที่ปรับปรุง')):
            mode = 'notes'; continue
        if mode == 'original':
            m['original'].append(text)
        elif mode == 'analysis':
            m['analysis'].append(text)
        elif mode == 'improved':
            m['paras'].append(text)
        elif mode == 'notes':
            m['notes'].append(text)
    t = m['topic'].replace('Topic:', '').strip()
    m['title'] = f'Improved Essay {k}: ' + (t[:80] + '…' if len(t) > 80 else t)
    models[f'IE{k}'] = m
if it:
    models['IE4']['table'] = [[c.text.strip() for c in r.cells] for r in it[0].rows]
if principles:
    models['IP'] = principles
models['IMG'] = {
    'title': 'ตัวอย่างลายมือ: Big salary is much more important than job satisfaction (291 คำ)', 'source': 'Image (1).pdf หน้า 1 (หน้า 2 ซ้ำ)',
    'topic': 'โจทย์: Big salary is much more important than job satisfaction.',
    'paras': [
        'It is often argued that it is more advantageous to choose a job with high wage even if it does not appeal to you at all. I completely disagree with this opinion and think that job satisfaction is much more important than salary.',
        'First of all, I believe that job satisfaction gives people a sense of fulfillment that no money can guarantee. Even though someone is earning a high salary, sometimes, they might feel tense. These people will not enjoy their life while pursuing one\u2019s interests will always bring pleasure and a feeling of satisfaction. For example, a lot of famous researchers made their career choices not because of appealing wages, but it is because they were passionate about science. That is why it is more important to choose the kind of work that makes you happy than to look only at a high salary.',
        'Secondly, doing what you like keeps you motivated and therefore leads to career growth. In other words, there is a strong relation between job satisfaction and productivity. People who love their jobs can easily excel in their fields of work and achieve better results than those who put a salary on the first place. For instance, my colleague, Henry Miller, decided to leave his everyday job despite a good wage and tried to become a writer. And after many enduring years of ups and downs, he became one of the most famous and well-paid authors of the twentieth first century. Thus, advantages of jobs that keep you satisfied are more important than the drawback of a low salary in a long-term perspective.',
        'In conclusion, I strongly believe that job satisfaction is more beneficial than a high salary because it makes people enjoy life and get inspire.'],
    'notes': ['โครงสร้างดี: paraphrase + จุดยืน → First of all (fulfilment) → Secondly (career growth) → In conclusion',
              'จุดที่ควรแก้: “twentieth first century” → “twenty-first century”; “get inspire” → “get inspired” / “feel inspired”; “on the first place” → “in first place”; “but it is because” → “but because”',
              'ตัวอย่าง “my colleague, Henry Miller … famous authors” ไม่สมเหตุสมผล (Henry Miller เป็นนักเขียนในอดีต) — ฉบับปรับปรุงใน Improved Essay 1 แก้เป็น Marie Curie / Henry Miller แบบอธิบายชัดเจน',
              'คำเชื่อมที่ใช้ดี: Even though, For example, That is why, In other words, For instance, Thus'],
}
for k2, v in LS.MODELS_SMALL.items():
    models[k2] = {'title': v['title'], 'source': 'ตัวอย่างในแอป', 'topic': '', 'paras': v['paras'], 'notes': v['notes']}

# ---------------- topics meta ----------------
TOPICS = {
    'R1': 'Aelt Writing 0421.pdf หน้า 1 (หัวข้อในเอกสาร)', 'R2': 'Aelt Writing 0421.pdf หน้า 1 — ระบุว่า “ข้อสอบจริง”',
    'R3': 'Aelt Writing 0421.pdf หน้า 1 — ระบุว่า “ข้อสอบจริง”', 'R4': 'Aelt Writing 0421.pdf หน้า 1 (หัวข้อในเอกสาร)',
    'R5': 'Aelt Writing 0421.pdf หน้า 1 (หัวข้อในเอกสาร)', 'R6': 'Aelt Writing 0421.pdf หน้า 1 (หัวข้อในเอกสาร)',
    'R7': 'Aelt Writing 0421.pdf หน้า 1 (มีตัวอย่างใน Image (1).pdf และ Improved Essay 1)', 'R8': 'Aelt Writing 0421.pdf หน้า 1 — “ข้อสอบจริง, 2020”',
    'R9': 'Aelt Writing 0421.pdf หน้า 2 — “ข้อสอบจริง, 2020”', 'R10': 'Aelt Writing 0421.pdf หน้า 2 — “ข้อสอบจริง, 2020”',
    'R11': 'Aelt Writing 0421.pdf หน้า 2 (หัวข้อในเอกสาร)', 'R13': 'Aelt Writing 0421.pdf หน้า 3 — ข้อสอบจริงรอบธันวาคม 2020',
    'R14': 'Aelt Writing 0421.pdf หน้า 3 — ข้อสอบจริงรอบมกราคม 2021', 'R15': 'Aelt Writing 0421.pdf หน้า 3 — ข้อสอบจริงเดือนกุมภาพันธ์ 2021',
    'R16': 'Aelt Writing 0421.pdf หน้า 3 — ข้อสอบจริงรอบมีนาคม 2021', 'RX': 'หัวข้อคาดการณ์ตาม “trend” (ไม่ใช่ข้อสอบจริง) — ใช้ฝึกรับมือหัวข้อใหม่',
}
topics = {k: {'label': v} for k, v in TOPICS.items()}

# ---------------- days ----------------
days = P42.build_days(vocab)
# the day each passage / test / grammar set belongs to follows the plan, not the content files
for d in days:
    for pid in (d.get('reading') or {}).get('passages', []):
        if pid in passages:
            passages[pid]['day'] = d['d']
    if d.get('test') in tests:
        tests[d['test']]['day'] = d['d']
        for pid in tests[d['test']]['passages']:
            passages[pid]['day'] = d['d']
    for gid in ([d['grammar']] if isinstance(d.get('grammar'), str) else d.get('grammar') or []):
        if gid in grammar:
            grammar[gid]['day'] = d['d']

# ---------------- guide extras ----------------
guide = dict(LS.GUIDE)
if guide_words:
    guide['words50'] = {'title': 'ศัพท์ 50 คำ (Guide)', 'blocks': [{'h': 'คำศัพท์สำคัญ 50 คำสำหรับ Opinion Essay (Opinion_Essay_Guide.docx ส่วนที่ 2)', 'table': guide_words}]}

# ---------------- validation ----------------
needed_p = {pid for d in days if d.get('reading') for pid in d['reading']['passages']}
for pid in sorted(needed_p - set(passages)):
    problems.append(f'passage {pid} missing')
for d in days:
    if d.get('test') and d['test'] not in tests:
        problems.append(f'test {d["test"]} missing')
    for gid in ([d['grammar']] if isinstance(d.get('grammar'), str) else d.get('grammar') or []):
        if gid not in grammar:
            problems.append(f'grammar {gid} missing')
    for t in [d.get('task', {}).get('topic'), (d.get('essay') or {}).get('topic')]:
        if t and t.startswith('R') and t not in essays:
            problems.append(f'essay {t} missing')
used_g = {g for d in days for g in ([d['grammar']] if isinstance(d.get('grammar'), str) else d.get('grammar') or [])}
used_p = needed_p | {pid for t in tests.values() for pid in t['passages']}
if used_g != set(grammar):
    problems.append(f'grammar sets not in plan: {sorted(set(grammar) - used_g)}')
if used_p != set(passages):
    problems.append(f'passages not in plan: {sorted(set(passages) - used_p)}')
for pid, p in passages.items():
    for i, q in enumerate(p['questions']):
        if q['type'] == 'mc':
            assert len(q['options']) == 4 and 0 <= q['answer'] <= 3, (pid, i)
        else:
            assert q.get('answers'), (pid, i)
for gid, g in grammar.items():
    for i, q in enumerate(g['quiz']):
        assert len(q['options']) == 4 and 0 <= q['answer'] <= 3, (gid, i)
# trim pass-through fields the app does not need
for p in passages.values():
    p.pop('source_note', None)

data = {
    'meta': {'start': P42.START.isoformat(), 'examDate': '2026-11-21', 'officialNov': '2026-11-14', 'built': datetime.date.today().isoformat(),
             'lightVocabDays': P42.LIGHT_VOCAB_DAYS, 'parts': P42.PARTS},
    'files': P.FILES, 'srcName': P.SRC_NAME, 'phases': P42.PHASES, 'days': days, 'vocab': vocab,
    'passages': passages, 'tests': tests, 'lessons': LS.RL, 'grammar': grammar, 'essays': essays,
    'topics': topics, 'models': models, 'blockInfo': LS.BLOCK_INFO, 'guide': guide,
}
js = '\n'.join((HERE / 'js' / f).read_text(encoding='utf-8') for f in ('core.js', 'quiz.js', 'day.js', 'views.js'))
html = (HERE / 'template.html').read_text(encoding='utf-8')
html = html.replace('__CSS__', (HERE / 'style.css').read_text(encoding='utf-8'))
html = html.replace('__DATA__', json.dumps(data, ensure_ascii=False, separators=(',', ':')).replace('<', '\\u003c'))
html = html.replace('__JS__', js)
OUT.parent.mkdir(parents=True, exist_ok=True)
OUT.write_text(html, encoding='utf-8')
(OUT.parent / 'data.json').write_text(json.dumps(data, ensure_ascii=False, indent=1), encoding='utf-8')
print('wrote', OUT, f'{OUT.stat().st_size/1e6:.2f} MB')
print('passages', len(passages), 'tests', len(tests), 'grammar', len(grammar), 'essays', len(essays), 'models', len(models), 'vocab', sum(len(v) for v in vocab.values()))
print('problems:', len(problems))
for x in problems:
    print('  -', x)
